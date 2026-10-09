"""dops_daily.py -- build the daily HOM1 dispatch sheet from DOPS + Cortex.

Inputs
  DOPS workbook  (daily "DOPS M D YYYY.xlsx" posted in Slack): sheets
                 Solution (route/wave/staging/packages), SD (scheduled stops),
                 CXL (pick-ups), MNR, EX (ignored, counted in Flags).
  Cortex roster  (Cortex > Operations > Delivery, today, HOM1): route code ->
                 driver. Accepts .json [{"name","route"}], .csv (name,route),
                 or .txt pasted page text (name / route code / phone repeating).

Output  Dispatch_MM-DD.xlsx : the dispatch sheet + a "Flags" sheet.
        Vehicle and Helper are left blank on purpose.

    python dops_daily.py --cortex cortex.txt                  # DOPS from Slack
    python dops_daily.py --dops "DOPS 10 9 2026.xlsx" --cortex cortex.json
    python dops_daily.py --date 2026-10-09 --cortex cortex.csv --out dops_out

Env / config (all optional except Slack when --dops is omitted):
    SLACK_BOT_TOKEN        bot token with channels:history (+groups:history for
                           a private channel) and files:read
    DOPS_SLACK_CHANNEL     channel id (or "slack_channel_id" in dops_config.json)
    DOPS_CONFIG            path to config json (default: dops_config.json beside
                           this file; gitignored -- put PIN/lock-box text there)
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import os
import re
import sys
from collections import OrderedDict, defaultdict
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

TZ = ZoneInfo("America/Chicago")
HERE = Path(__file__).resolve().parent
CONFIG_PATH = Path(os.environ.get("DOPS_CONFIG", HERE / "dops_config.json"))

DEFAULTS = {
    "station": "HOM 1",
    "dispatcher": "Ed",
    "slack_channel_id": "",
    # dock per staging slot ("A.3" -> docks["A"][2]); None = no dock
    "docks": {"A": [111, 112, 113, 114, 115],
              "B": [111, 112, 113, 114, 115],
              "C": [111, None]},
    "rts_hours_after_wave": 10,
    "otd_minutes_after_wave": 35,
    "heavy_far_note": "Heavy/Far routes Go for the next open dock!",
    "footer_legend": "A= AMAZON   P=Cargo   C=Rental",
    "footer_phone": "",          # e.g. "PHONE PIN ____/LOCK BOX ____" -- keep in the gitignored config
    "footer_otd": "On time delivery (OTD)",
    "meeting_note": "",          # e.g. "Meeting in Break room @7:15AM Sharp"
    "name_overrides": {},        # {"Full Name From Cortex": "Display"}
}


def load_config() -> dict:
    cfg = json.loads(json.dumps(DEFAULTS))
    if CONFIG_PATH.exists():
        cfg.update(json.loads(CONFIG_PATH.read_text(encoding="utf-8")))
    return cfg


# ─── DOPS workbook ─────────────────────────────────────────────────────────

def _rows(wb, name: str):
    if name not in wb.sheetnames:
        return []
    return [r for r in wb[name].iter_rows(min_row=2, values_only=True) if r and len(r) > 1 and r[1]]


def _parse_wave(v) -> dt.time:
    if isinstance(v, dt.datetime):
        return v.time()
    if isinstance(v, dt.time):
        return v
    return dt.datetime.strptime(str(v).strip().upper(), "%I:%M %p").time()


def read_dops(path: str | Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    routes = []                                   # Solution, in sheet order
    for r in _rows(wb, "Solution"):
        routes.append({"route": str(r[1]).strip(), "wave": _parse_wave(r[3]),
                       "staging": str(r[4]).strip(), "pkgs": int(r[5] or 0)})
    sd, cxl, mnr = (OrderedDict() for _ in range(3))
    for r in _rows(wb, "SD"):
        sd.setdefault(str(r[1]).strip(), []).append(
            {"tba": r[0], "start": r[2], "end": r[3], "service": r[4] or "", "city": r[5] or "", "order": r[6]})
    for r in _rows(wb, "CXL"):
        cxl.setdefault(str(r[1]).strip(), []).append(r[0])
    for r in _rows(wb, "MNR"):
        mnr.setdefault(str(r[1]).strip(), []).append(r[0])
    ex = _rows(wb, "EX")
    return {"routes": routes, "sd": sd, "cxl": cxl, "mnr": mnr, "ex_count": len(ex)}


# ─── Cortex roster ─────────────────────────────────────────────────────────

_ROUTE_RE = re.compile(r"^[A-Z]{1,3}\d{1,3}$")


def read_cortex(path: str | Path) -> list[tuple[str, str]]:
    """-> [(full_name, route)]"""
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".json":
        return [(d["name"].strip(), d["route"].strip().upper()) for d in json.loads(p.read_text(encoding="utf-8"))]
    if ext == ".csv":
        with p.open(newline="", encoding="utf-8-sig") as f:
            rd = csv.DictReader(f)
            low = {k.lower().strip(): k for k in rd.fieldnames or []}
            return [(row[low["name"]].strip(), row[low["route"]].strip().upper()) for row in rd]
    lines = [l.strip() for l in p.read_text(encoding="utf-8").splitlines() if l.strip()]
    out = []
    for i, l in enumerate(lines):
        if _ROUTE_RE.match(l) and i > 0 and not _ROUTE_RE.match(lines[i - 1]):
            out.append((lines[i - 1], l))
    return out


# ─── Slack ─────────────────────────────────────────────────────────────────

def fetch_dops_from_slack(day: dt.date, dest: Path, cfg: dict) -> Path:
    import requests
    token = os.environ["SLACK_BOT_TOKEN"]
    channel = os.environ.get("DOPS_SLACK_CHANNEL") or cfg["slack_channel_id"]
    if not channel:
        raise SystemExit("No Slack channel id: set DOPS_SLACK_CHANNEL or slack_channel_id in dops_config.json")
    hdr = {"Authorization": f"Bearer {token}"}
    oldest = dt.datetime(day.year, day.month, day.day, tzinfo=TZ).timestamp()
    r = requests.get("https://slack.com/api/conversations.history", headers=hdr,
                     params={"channel": channel, "oldest": oldest, "limit": 200}, timeout=30).json()
    if not r.get("ok"):
        raise SystemExit(f"Slack error: {r.get('error')}")
    want = re.compile(rf"DOPS[\s_-]*{day.month}[\s_-]*{day.day}[\s_-]*{day.year}", re.I)
    for m in r["messages"]:                       # newest first -> latest re-post wins
        for f in m.get("files", []):
            if want.search(f.get("name", "")) and f["name"].lower().endswith(".xlsx"):
                data = requests.get(f["url_private_download"], headers=hdr, timeout=60)
                data.raise_for_status()
                dest.mkdir(parents=True, exist_ok=True)
                out = dest / f["name"]
                out.write_bytes(data.content)
                return out
    raise SystemExit(f"No DOPS {day.month} {day.day} {day.year} xlsx found in channel today")


# ─── Matching / flags ──────────────────────────────────────────────────────

def display_names(assigned: dict[str, str], overrides: dict) -> dict[str, str]:
    """route -> display name. First name; first + last initial on collision."""
    firsts = {}
    for route, full in assigned.items():
        firsts[route] = overrides.get(full) or full.split()[0]
    counts = defaultdict(int)
    for n in firsts.values():
        counts[n] += 1
    out = {}
    for route, full in assigned.items():
        n = firsts[route]
        parts = full.split()
        out[route] = f"{n} {parts[-1][0]}." if counts[n] > 1 and len(parts) > 1 and not overrides.get(full) else n
    return out


def _hr12(d: dt.datetime) -> int:
    return d.hour % 12 or 12


def _fmt_time(t: dt.time) -> str:
    return f"{t.hour % 12 or 12}:{t.minute:02d}"


# ─── Sheet ─────────────────────────────────────────────────────────────────

WHITE = "FFFFFF"
FILL = {k: PatternFill("solid", fgColor=v) for k, v in {
    "gray": "D9D9D9", "orange": "FF9900", "red": "FF0000", "green": "6AA84F", "pink": "D5A6BD",
    "teal": "26C6DA", "ember": "F08A4B", "lgreen": "B6D7A8", "lpink": "EAD1DC",
    "yellow": "FFFF00", "warn": "FFE599", "white": WHITE}.items()}
_THIN = Side(style="thin", color="999999")
BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")


def _put(ws, cell, value=None, *, fill=None, bold=False, size=10, color="000000", align=CENTER, border=True):
    c = ws[cell]
    if value is not None:
        c.value = value
    if fill:
        c.fill = FILL[fill]
    c.font = Font(bold=bold, size=size, color=color)
    c.alignment = align
    if border:
        c.border = BOX
    return c


def _merge(ws, rng: str, value=None, **kw):
    first = rng.split(":")[0]
    ws.merge_cells(rng)
    for row in ws[rng]:
        for c in row:
            c.border = BOX
            if kw.get("fill"):
                c.fill = FILL[kw["fill"]]
    return _put(ws, first, value, **kw)


def build(dops: dict, roster: list[tuple[str, str]], day: dt.date, cfg: dict, out_path: Path) -> list[str]:
    flags: list[str] = []
    routes = dops["routes"]
    dops_codes = {r["route"] for r in routes}

    # route -> full name from Cortex
    by_route: dict[str, str] = {}
    for full, route in roster:
        if route in by_route and by_route[route] != full:
            flags.append(f"Cortex lists route {route} twice ({by_route[route]} / {full}); using {by_route[route]}")
            continue
        by_route.setdefault(route, full)
    assigned = {r["route"]: by_route[r["route"]] for r in routes if r["route"] in by_route}
    names = display_names(assigned, cfg["name_overrides"])
    for r in routes:
        if r["route"] not in by_route:
            flags.append(f"No Cortex driver for route {r['route']} ({r['staging']}) -- Driver cell left as '?? {r['route']}'")
    for route, full in by_route.items():
        if route not in dops_codes:
            flags.append(f"Cortex has {full} on {route}, which is not in DOPS Solution -- not placed")
    for route, n in names.items():
        if n != assigned[route].split()[0] and not cfg["name_overrides"].get(assigned[route]):
            flags.append(f"Two drivers share first name {assigned[route].split()[0]!r}; shown as {n!r}")

    def who(route: str) -> str:
        return names.get(route, route)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = day.strftime("%m-%d")
    widths = dict(A=7, B=11, C=15, D=11, E=9, F=13, G=8, H=8, I=11,
                  J=18, K=12, L=9, M=7, N=34, O=17, P=15, Q=12, R=8, S=18, T=12)
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    # header band
    _merge(ws, "B1:E2", cfg["station"], fill="gray", bold=True, size=28, align=LEFT)
    _merge(ws, "F1:I2", f"Date: {day.month}/{day.day}", fill="gray", bold=True, size=12, align=LEFT)
    _merge(ws, "B3:G3", f"{day.strftime('%A')} Dispatch: {cfg['dispatcher']}", fill="gray", bold=True, size=11)
    _merge(ws, "H3:I3", "Dispatch Only", fill="gray", bold=True)

    # ── left: waves ──
    waves: "OrderedDict[dt.time, list]" = OrderedDict()
    for r in sorted(routes, key=lambda x: (x["wave"], x["staging"])):
        waves.setdefault(r["wave"], []).append(r)
    total_pkgs = sum(r["pkgs"] for r in routes)
    row, bag, first_wave = 4, 1, True
    for i, (wtime, wroutes) in enumerate(waves.items(), start=1):
        wdt = dt.datetime.combine(day, wtime)
        rts = wdt + dt.timedelta(hours=cfg["rts_hours_after_wave"])
        otd = wdt + dt.timedelta(minutes=cfg["otd_minutes_after_wave"])
        _merge(ws, f"B{row}:G{row}", f"WAVE {i}   {_fmt_time(wtime)} {'AM' if wtime.hour < 12 else 'PM'}      RTS: {_fmt_time(rts.time())}",
               fill="orange", bold=True, size=11)
        if first_wave:
            _put(ws, f"H{row}", "PKG Count", fill="gray", bold=True, size=9)
            _put(ws, f"I{row}", total_pkgs, fill="white", bold=True)
        else:
            _put(ws, f"H{row}", fill="white"); _put(ws, f"I{row}", fill="white")
        row += 1
        for col, h in zip("BCDEFG", ["Staging", "Driver", "Vehicle", "Route", "Helper", "Dock"]):
            _put(ws, f"{col}{row}", h, fill="orange", bold=True, size=9)
        _put(ws, f"H{row}", "Bag" if first_wave else None, fill="gray" if first_wave else "white", bold=True, size=9)
        _put(ws, f"I{row}", "5 Star Team" if first_wave else None, fill="gray" if first_wave else "white", bold=True, size=9)
        first_wave = False
        row += 1
        for r in wroutes:
            letter, _, idx = r["staging"].partition(".")
            try:
                dock = cfg["docks"][letter][int(idx) - 1]
            except (KeyError, IndexError, ValueError):
                dock = None
                flags.append(f"No dock configured for staging {r['staging']}")
            have = r["route"] in names
            _put(ws, f"B{row}", r["staging"], bold=True)
            _put(ws, f"C{row}", names.get(r["route"], f"?? {r['route']}"), bold=True, fill=None if have else "warn")
            _put(ws, f"D{row}", fill="gray")                      # Vehicle -- blank by design
            _put(ws, f"E{row}", r["route"], bold=True)
            _put(ws, f"F{row}")                                   # Helper -- blank by design
            _put(ws, f"G{row}", dock, fill=None if dock else "gray")
            _put(ws, f"H{row}", bag)
            _put(ws, f"I{row}")
            bag += 1
            row += 1
        _merge(ws, f"B{row}:F{row}", cfg["heavy_far_note"], bold=True, size=10)
        _put(ws, f"G{row}", f"{_fmt_time(otd.time())} OTD", fill="yellow", bold=True, size=9)
        _put(ws, f"H{row}", fill="white"); _put(ws, f"I{row}", fill="white")
        row += 1
    row += 1
    for txt, fill, rng in [(cfg["footer_legend"], "yellow", f"B{row}:C{row}"),
                           (cfg["footer_phone"], "yellow", f"D{row}:F{row}"),
                           (cfg["footer_otd"], "yellow", f"G{row}:I{row}")]:
        _merge(ws, rng, txt or None, fill=fill, bold=True, size=8)
    if cfg["meeting_note"]:
        row += 1
        _merge(ws, f"B{row}:I{row}", cfg["meeting_note"], fill="yellow", bold=True, size=11)

    # ── right: scheduled stops / pick ups / MNR ──
    _merge(ws, "J1:O2", "Scheduled Stops", fill="red", bold=True, size=18)
    _merge(ws, "P1:R2", "Pick Ups", fill="green", bold=True, size=18)
    _merge(ws, "S1:T2", "MNR", fill="pink", bold=True, size=18)
    for col, h, f in [("J", "TBA", "red"), ("K", "Driver", "red"), ("L", "Driver Aid #", "red"), ("M", "Time", "red"),
                      ("N", "Service", "red"), ("O", "City", "red"), ("P", "CRXL", "green"), ("Q", "Route", "green"),
                      ("R", "Time", "green"), ("S", "TBA", "pink"), ("T", "Route", "pink")]:
        _put(ws, f"{col}3", h, fill=f, bold=True, size=9)

    r_sd = 4
    for gi, (route, stops) in enumerate(dops["sd"].items()):
        band = "teal" if gi % 2 == 0 else "ember"
        for s in stops:
            win = f"{_hr12(s['start'])}-{_hr12(s['end'])}" if s["start"] and s["end"] else ""
            _put(ws, f"J{r_sd}", s["tba"], fill=band, size=8, align=LEFT)
            _put(ws, f"K{r_sd}", who(route), fill=band, size=8)
            _put(ws, f"L{r_sd}", s["order"], fill=band, size=8)
            _put(ws, f"M{r_sd}", win, fill="red", bold=True, size=8, color=WHITE)
            _put(ws, f"N{r_sd}", s["service"], fill=band, size=8, align=LEFT)
            _put(ws, f"O{r_sd}", s["city"], fill=band, size=8, align=LEFT)
            r_sd += 1
        r_sd += 1
        if route not in dops_codes:
            flags.append(f"SD has stops for {route}, which is not in DOPS Solution")
    r_cx = 4
    for route, ids in dops["cxl"].items():
        for t in ids:
            _put(ws, f"P{r_cx}", t, fill="lgreen", size=8, align=LEFT)
            _put(ws, f"Q{r_cx}", who(route), fill="lgreen", size=8)
            _put(ws, f"R{r_cx}", fill="lgreen")                    # pickup window is not in DOPS
            r_cx += 1
        r_cx += 1
    r_mn = 4
    sd_ids = {s["tba"] for stops in dops["sd"].values() for s in stops}
    for route, ids in dops["mnr"].items():
        for t in ids:
            _put(ws, f"S{r_mn}", t, fill="lpink", size=8, align=LEFT)
            _put(ws, f"T{r_mn}", who(route), fill="lpink", size=8)
            if t in sd_ids:
                flags.append(f"{t} appears in both SD and MNR ({route})")
            r_mn += 1
        r_mn += 1
    if dops["cxl"]:
        flags.append("Pick Ups 'Time' column left blank (pickup windows are not in the DOPS workbook)")
    if dops["ex_count"]:
        flags.append(f"EX sheet has {dops['ex_count']} package(s); not placed on the dispatch sheet")
    flags.append(f"PKG Count shows total Num Packages from DOPS Solution ({total_pkgs})")

    last = max(row, r_sd, r_cx, r_mn) - 1
    _merge(ws, f"A1:A{last}", "M A Z E", fill="gray", bold=True, size=22)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="top", text_rotation=255, wrap_text=True)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)

    fs = wb.create_sheet("Flags")
    fs["A1"] = f"Flags for {day.isoformat()}"
    fs["A1"].font = Font(bold=True, size=12)
    for i, f in enumerate(flags, start=3):
        fs[f"A{i}"] = f
    fs.column_dimensions["A"].width = 120
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return flags


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dops", help="DOPS xlsx; omit to fetch today's from Slack")
    ap.add_argument("--cortex", required=True, help="Cortex roster (.json/.csv/.txt)")
    ap.add_argument("--date", help="YYYY-MM-DD (default: today, America/Chicago)")
    ap.add_argument("--out", default="dops_out", help="output directory")
    a = ap.parse_args(argv)

    cfg = load_config()
    day = dt.date.fromisoformat(a.date) if a.date else dt.datetime.now(TZ).date()
    out_dir = Path(a.out)
    dops_path = Path(a.dops) if a.dops else fetch_dops_from_slack(day, out_dir / "inbox", cfg)
    dops = read_dops(dops_path)
    roster = read_cortex(a.cortex)
    out = out_dir / f"Dispatch_{day.strftime('%m-%d')}.xlsx"
    flags = build(dops, roster, day, cfg, out)
    print(f"wrote {out}  ({len(dops['routes'])} routes, {len(roster)} Cortex drivers)")
    for f in flags:
        print(f"  ! {f}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
